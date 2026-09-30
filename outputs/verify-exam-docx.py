# -*- coding: utf-8 -*-
"""校验生成的 docx 与题库 JSON 是否逐条一致
用法：C:\\Users\\Administrator\\.workbuddy\\binaries\\python\\envs\\html2docx\\Scripts\\python.exe outputs/verify-exam-docx.py
"""
import json, io, sys, os, re
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

ROOT = r'C:\Users\Administrator\WorkBuddy\2026-07-28-10-50-05'
quiz = json.load(io.open(os.path.join(ROOT, 'outputs', 'miyang-quiz.json'), encoding='utf-8'))
doc = Document(os.path.join(ROOT, 'outputs', '弥生花色价格考核试卷_含答案卷.docx'))

N = lambda s: re.sub(r'\s+', '', s)
paras = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
full = N('\n'.join(paras))
ai = full.find(N('答案卷 · 主管用'))
tail = full[ai:] if ai >= 0 else ''
L = lambda i: chr(65 + i)

ok = fail = 0
def chk(name, cond, extra=''):
    global ok, fail
    if cond:
        ok += 1
    else:
        fail += 1
        print('FAIL', name, extra)

chk('标题在卷', N(quiz['meta']['title']) in full)
chk('答案卷分节存在', ai >= 0)
chk('题量说明', N('共 %d 题' % quiz['meta']['counts']) in full)
chk('满分说明', N('%d 分（80 分合格）' % quiz['meta']['total']) in full)

for q in quiz['questions']:
    chk('Q%d 题干' % q['no'], N(q['stem'])[:12] in full)
    if q.get('options'):
        chk('Q%d 选项A' % q['no'], N(q['options'][0])[:8] in full)
    if q['type'] == 'match':
        a = N(q['pairs'][0]['right'])
    elif q['type'] == 'text':
        a = N(q['answer'])[:10]
    elif isinstance(q['answer'], list):
        a = '、'.join(L(i) for i in q['answer'])
    else:
        a = L(q['answer']) if q['type'] == 'single' else ('正确' if q['answer'] == 0 else '错误')
    chk('Q%d 答案在答案卷' % q['no'], a in tail, a[:20])

chk('解析已写入', '解析：' in full)
chk('配对表存在', len(doc.tables) >= 2)
print('=== docx 校验: pass=%d fail=%d | 段落=%d 表格=%d ===' % (ok, fail, len(doc.paragraphs), len(doc.tables)))
sys.exit(1 if fail else 0)
